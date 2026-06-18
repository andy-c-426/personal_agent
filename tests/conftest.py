import pytest
import tempfile
import shutil
from pathlib import Path


@pytest.fixture
def temp_dir():
    d = tempfile.mkdtemp()
    yield Path(d)
    shutil.rmtree(d)


@pytest.fixture
def sample_md_file(temp_dir):
    p = temp_dir / "test.md"
    p.write_text("# Test Doc\n\nThis is a test document for knowledge base ingestion.")
    return p


@pytest.fixture
def sample_pdf_file(temp_dir):
    p = temp_dir / "test.pdf"
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, text="Test PDF content for ingestion pipeline")
    pdf.output(str(p))
    return p


@pytest.fixture
def sample_docx_file(temp_dir):
    from docx import Document
    p = temp_dir / "test.docx"
    doc = Document()
    doc.add_paragraph("Test Word document content for ingestion pipeline")
    doc.save(str(p))
    return p


@pytest.fixture
def sample_text_file(temp_dir):
    p = temp_dir / "test.txt"
    p.write_text("Plain text content for the knowledge base.")
    return p


@pytest.fixture
def quality_docs(temp_dir):
    """Hand-crafted documents for RAG quality evaluation."""
    docs = {
        "python_intro.md": (
            "# Python Programming\n\n"
            "Python is a high-level, interpreted programming language known for "
            "its readability and versatility. It was created by Guido van Rossum "
            "and first released in 1991. Python supports multiple programming "
            "paradigms including procedural, object-oriented, and functional programming.\n\n"
            "## Key Features\n\n"
            "Python's design philosophy emphasizes code readability through its "
            "use of significant indentation. Its language constructs and "
            "object-oriented approach aim to help programmers write clear, "
            "logical code for small and large-scale projects."
        ),
        "machine_learning.md": (
            "# Machine Learning Fundamentals\n\n"
            "Machine learning is a subset of artificial intelligence that enables "
            "systems to learn and improve from experience without being explicitly "
            "programmed. The core idea is to develop algorithms that can receive "
            "input data and use statistical analysis to predict an output.\n\n"
            "## Types of Machine Learning\n\n"
            "Supervised learning uses labeled training data to learn a mapping "
            "from inputs to outputs. Unsupervised learning finds hidden patterns "
            "in unlabeled data. Reinforcement learning trains agents to make "
            "sequences of decisions through reward signals.\n\n"
            "## Deep Learning\n\n"
            "Deep learning is a specialized form of machine learning that uses "
            "neural networks with many layers (hence 'deep') to progressively "
            "extract higher-level features from raw input."
        ),
        "cooking_tips.md": (
            "# Essential Cooking Tips\n\n"
            "Good cooking starts with fresh ingredients. Always read the entire "
            "recipe before you begin cooking. Mise en place — having all your "
            "ingredients prepared and measured before you start — is the "
            "foundation of efficient cooking.\n\n"
            "## Seasoning\n\n"
            "Salt is the most important seasoning in any kitchen. Add salt in "
            "layers throughout the cooking process rather than all at once at "
            "the end. Taste as you go and adjust seasoning accordingly."
        ),
        "climate_science.md": (
            "# Climate Change Overview\n\n"
            "Climate change refers to long-term shifts in temperatures and "
            "weather patterns. These shifts may be natural, but since the 1800s, "
            "human activities have been the main driver of climate change, "
            "primarily due to the burning of fossil fuels.\n\n"
            "## Greenhouse Effect\n\n"
            "The greenhouse effect is the process through which heat is trapped "
            "near Earth's surface by greenhouse gases. Carbon dioxide, methane, "
            "and water vapor are the primary greenhouse gases."
        ),
        "product_review.md": (
            "# Smartphone Review: Model X Pro\n\n"
            "The Model X Pro features a 6.7-inch OLED display with 120Hz refresh "
            "rate. Battery life is excellent at 5000mAh, lasting a full day of "
            "heavy use. The camera system includes a 108MP main sensor, 12MP "
            "ultrawide, and 10MP telephoto with 3x optical zoom.\n\n"
            "## Performance\n\n"
            "Powered by the latest Snapdragon processor and 12GB of RAM, the "
            "Model X Pro handles multitasking and gaming with ease. Storage "
            "options include 128GB, 256GB, and 512GB variants."
        ),
    }
    for name, content in docs.items():
        p = temp_dir / name
        p.write_text(content)
    return temp_dir
